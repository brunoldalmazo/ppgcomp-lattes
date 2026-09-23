from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.config import (
    OUTPUT_DIR,
    REFERENCE_YEAR,
    PERIOD_START,
    PERIOD_YEARS,
    MIN_SCORE,
    LATTES_UPDATE_LIMIT_DAYS,
    QUALIS_AREA,
    QUALIS_FILE,
    WEIGHTS,
)


def autosize(ws):
    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column

        for cell in column_cells:
            if cell.value is None:
                continue

            max_length = max(
                max_length,
                len(str(cell.value)),
            )

        ws.column_dimensions[
            get_column_letter(column)
        ].width = min(
            max_length + 2,
            80,
        )


def style_header(ws):
    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.row_dimensions[1].height = 30


def style_body(ws):
    for row in ws.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def add_sheet(
    wb,
    name,
    headers,
    rows,
):
    ws = wb.create_sheet(name)

    ws.append(headers)

    for row in rows:
        ws.append(row)

    style_header(ws)
    style_body(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    return ws


def export_excel(resultado):
    year_output_dir = (
        OUTPUT_DIR
        / str(REFERENCE_YEAR)
    )

    year_output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    output_file = (
        year_output_dir
        / f"avaliacao_ppgcomp_{REFERENCE_YEAR}.xlsx"
    )

    wb = Workbook()

    # ============================================================
    # RESUMO
    # ============================================================

    ws = wb.active
    ws.title = "Resumo"

    headers = [
        "Professor",
        "Última atualização Lattes",
        "A1 qtd.",
        "A1 pontos",
        "A2 qtd.",
        "A2 pontos",
        "A3 qtd.",
        "A3 pontos",
        "A4 qtd.",
        "A4 pontos",
        "B1 qtd.",
        "B1 pontos",
        "B2 qtd.",
        "B2 pontos",
        "B3 qtd.",
        "B3 pontos",
        "B4 qtd.",
        "B4 pontos",
        "C qtd.",
        "C pontos",
        "Total",
        "Status",
        "Publicações",
        "Publicações em coautoria",
    ]

    ws.append(headers)

    for professor in sorted(
        resultado["professors"]
    ):
        dados = resultado["professors"][
            professor
        ]

        counts = dados.get(
            "counts",
            {},
        )

        points = dados.get(
            "points",
            {},
        )

        ws.append([
            professor,
            dados.get(
                "lattes_update",
                "",
            ),
            counts.get("A1", 0),
            points.get("A1", 0),
            counts.get("A2", 0),
            points.get("A2", 0),
            counts.get("A3", 0),
            points.get("A3", 0),
            counts.get("A4", 0),
            points.get("A4", 0),
            counts.get("B1", 0),
            points.get("B1", 0),
            counts.get("B2", 0),
            points.get("B2", 0),
            counts.get("B3", 0),
            points.get("B3", 0),
            counts.get("B4", 0),
            points.get("B4", 0),
            counts.get("C", 0),
            points.get("C", 0),
            dados.get("total", 0),
            dados.get("status", ""),
            dados.get("publications", 0),
            dados.get(
                "coauthored_publications",
                0,
            ),
        ])

    style_header(ws)
    style_body(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    # ============================================================
    # PRODUÇÃO
    # ============================================================

    production_rows = []

    for publication in resultado[
        "publications"
    ]:
        evaluated_professors = publication.get(
            "evaluated_professors",
            [],
        )

        production_rows.append([
            publication.get(
                "year",
                "",
            ),
            publication.get(
                "type",
                "",
            ),
            publication.get(
                "title",
                "",
            ),
            publication.get(
                "authors",
                "",
            ),
            publication.get(
                "venue",
                "",
            ),
            publication.get(
                "issn",
                "",
            ),
            publication.get(
                "doi",
                "",
            ),
            publication.get(
                "qualis",
                "",
            ),
            publication.get(
                "qualis_points",
                0,
            ),
            publication.get(
                "qualis_title",
                "",
            ),
            publication.get(
                "qualis_issn",
                "",
            ),
            publication.get(
                "qualis_fonte",
                "",
            ),
            publication.get(
                "qualis_periodo",
                "",
            ),
            publication.get(
                "qualis_matched_by",
                "",
            ),
            ", ".join(
                evaluated_professors
            ),
            publication.get(
                "evaluated_coauthor_count",
                0,
            ),
            publication.get(
                "individual_points",
                0,
            ),
        ])

    add_sheet(
        wb,
        "Produção",
        [
            "Ano",
            "Tipo",
            "Título",
            "Autores",
            "Veículo",
            "ISSN",
            "DOI",
            "Qualis",
            "Pontos Qualis",
            "Título Qualis",
            "ISSN Qualis",
            "Fonte Qualis",
            "Período Qualis",
            "Tipo de match Qualis",
            "Professores avaliados",
            "Nº professores avaliados",
            "Pontos individuais",
        ],
        production_rows,
    )

    # ============================================================
    # COAUTORIA - DETALHAMENTO
    # ============================================================

    coauthorship_rows = []

    for professor in sorted(
        resultado["coauthorships"]
    ):
        publications = resultado[
            "coauthorships"
        ][professor]

        for publication in publications:
            evaluated = publication.get(
                "evaluated_professors",
                [],
            )

            coauthors = [
                name
                for name in evaluated
                if name != professor
            ]

            coauthorship_rows.append([
                professor,
                ", ".join(coauthors),
                publication.get(
                    "year",
                    "",
                ),
                publication.get(
                    "type",
                    "",
                ),
                publication.get(
                    "title",
                    "",
                ),
                publication.get(
                    "venue",
                    "",
                ),
                publication.get(
                    "issn",
                    "",
                ),
                publication.get(
                    "doi",
                    "",
                ),
                publication.get(
                    "qualis",
                    "",
                ),
                publication.get(
                    "qualis_points",
                    0,
                ),
                publication.get(
                    "individual_points",
                    0,
                ),
            ])

    add_sheet(
        wb,
        "Coautoria",
        [
            "Professor",
            "Coautores avaliados",
            "Ano",
            "Tipo",
            "Título",
            "Veículo",
            "ISSN",
            "DOI",
            "Qualis",
            "Pontos Qualis",
            "Pontos individuais",
        ],
        coauthorship_rows,
    )

    # ============================================================
    # COAUTORIA - RESUMO POR PAR
    # ============================================================

    pair_counts = defaultdict(int)

    for professor in resultado[
        "coauthorships"
    ]:
        for publication in resultado[
            "coauthorships"
        ][professor]:

            evaluated = sorted(
                set(
                    publication.get(
                        "evaluated_professors",
                        [],
                    )
                )
            )

            for i in range(
                len(evaluated)
            ):
                for j in range(
                    i + 1,
                    len(evaluated),
                ):
                    pair = (
                        evaluated[i],
                        evaluated[j],
                    )

                    pair_counts[pair] += 1

    pair_rows = []

    for (
        professor_a,
        professor_b,
    ), count in sorted(
        pair_counts.items()
    ):
        pair_rows.append([
            professor_a,
            professor_b,
            count,
        ])

    add_sheet(
        wb,
        "Pares de coautoria",
        [
            "Professor A",
            "Professor B",
            "Publicações em conjunto",
        ],
        pair_rows,
    )

    # ============================================================
    # PENDÊNCIAS
    # ============================================================

    pending_rows = []

    for pending in resultado[
        "pending"
    ]:
        pending_rows.append([
            pending.get(
                "professor",
                "",
            ),
            pending.get(
                "year",
                "",
            ),
            pending.get(
                "type",
                "",
            ),
            pending.get(
                "title",
                "",
            ),
            pending.get(
                "venue",
                "",
            ),
            pending.get(
                "issn",
                "",
            ),
            pending.get(
                "doi",
                "",
            ),
            pending.get(
                "pendencia",
                pending.get(
                    "reason",
                    "",
                ),
            ),
        ])

    add_sheet(
        wb,
        "Pendências",
        [
            "Professor",
            "Ano",
            "Tipo",
            "Título",
            "Veículo",
            "ISSN",
            "DOI",
            "Motivo",
        ],
        pending_rows,
    )

    # ============================================================
    # CONFIGURAÇÃO
    # ============================================================

    configuration_rows = [
        [
            "Ano de referência",
            REFERENCE_YEAR,
        ],
        [
            "Período em anos",
            PERIOD_YEARS,
        ],
        [
            "Início do período",
            PERIOD_START,
        ],
        [
            "Pontuação mínima",
            MIN_SCORE,
        ],
        [
            "Limite atualização Lattes (dias)",
            LATTES_UPDATE_LIMIT_DAYS,
        ],
        [
            "Área Qualis",
            QUALIS_AREA,
        ],
        [
            "Arquivo Qualis",
            str(QUALIS_FILE),
        ],
    ]

    for qualis, points in WEIGHTS.items():
        configuration_rows.append([
            f"Peso {qualis}",
            points,
        ])

    add_sheet(
        wb,
        "Configuração",
        [
            "Parâmetro",
            "Valor",
        ],
        configuration_rows,
    )

    # ============================================================
    # ÍNDICE DE PUBLICAÇÕES
    # ============================================================

    index_rows = []

    for key, professors in sorted(
        resultado[
            "publication_index"
        ].items(),
        key=lambda item: str(item[0]),
    ):
        index_rows.append([
            str(key),
            ", ".join(professors),
        ])

    add_sheet(
        wb,
        "Índice publicações",
        [
            "Chave da publicação",
            "Professores",
        ],
        index_rows,
    )

    wb.save(output_file)

    return output_file
